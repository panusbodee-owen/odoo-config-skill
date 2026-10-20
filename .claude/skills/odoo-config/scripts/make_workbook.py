#!/usr/bin/env python3
"""Generate a configuration workbook (customer decision log) as Excel.

Usage: make_workbook.py OUT.xlsx --version 18 [--modules general,accounting,sales,...]
       [--project NAME]
Sheets: Cover, one per module (settings decisions), master lists (Taxes, Journals,
Sequences, Warehouses, Approval Matrix, Roles), Sign-off.
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from items import LISTS, SETTINGS  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

FONT = "Tahoma"
HEAD = PatternFill("solid", fgColor="714B67")
thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LIST_FOR = {"accounting": ["Taxes", "Journals", "Sequences"], "inventory": ["Warehouses"],
            "purchase": ["Approval Matrix"], "sales": ["Approval Matrix"], "general": ["Roles"]}


def style(ws, widths):
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row == 1, color="FFFFFF" if c.row == 1 else "000000")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
            if c.row == 1:
                c.fill = HEAD
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


def main():
    p = argparse.ArgumentParser()
    p.add_argument("out")
    p.add_argument("--version", type=int, default=18)
    p.add_argument("--modules", default="general,accounting,sales,purchase,inventory")
    p.add_argument("--project", default="")
    a = p.parse_args()
    mods = [m.strip() for m in a.modules.split(",") if m.strip()]
    bad = [m for m in mods if m not in SETTINGS]
    if bad:
        sys.exit(f"unknown modules: {bad}; known: {list(SETTINGS)}")

    wb = Workbook()
    cv = wb.active
    cv.title = "Cover"
    for r in (["Configuration Workbook", a.project], ["Odoo version", a.version],
              ["Modules", ", ".join(mods)], ["Status values", "Open / Decided / Configured / Verified"],
              ["Note", "Setting names follow Odoo 18 UI; names marked (verify) must be checked on the target version."]):
        cv.append(r)
    cv.column_dimensions["A"].width, cv.column_dimensions["B"].width = 22, 90

    status = '"Open,Decided,Configured,Verified"'
    lists = []
    for m in mods:
        ws = wb.create_sheet(m.capitalize())
        ws.append(["ID", "Area", "Setting", "Where in Odoo", "Options", "Recommended",
                   "Customer decision", "Owner", "Status", "Remark"])
        for i, (area, setting, where, opts, rec, note) in enumerate(SETTINGS[m], 1):
            ws.append([f"CFG-{m[:3].upper()}-{i:02d}", area, setting, where, opts, rec, "", "", "Open", note])
        dv = DataValidation(type="list", formula1=status, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"I2:I{ws.max_row + 50}")
        style(ws, (13, 14, 34, 34, 26, 20, 26, 14, 12, 34))
        for name in LIST_FOR.get(m, []):
            if name not in lists:
                lists.append(name)
    for name in lists:
        cols, rows = LISTS[name]
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in rows:
            ws.append(r)
        style(ws, [22] * len(cols))
    so = wb.create_sheet("Sign-off")
    so.append(["Role", "Name", "Signature", "Date"])
    for r in ("Customer Process Owner", "Customer Project Manager", "Implementer Consultant", "Implementer Project Manager"):
        so.append([r, "", "", ""])
    style(so, (32, 28, 24, 16))
    wb.save(a.out)
    print(a.out, "sheets:", ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
